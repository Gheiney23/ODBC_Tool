import pandas as pd
import pyodbc
import sqlalchemy as sa
from openpyxl import load_workbook

# Creating DB connection
connect_info = ('DRIVER=driver_here;'
                 'SERVER=server_here;'
                 'DATABASE=Database_here;'
                 'Trusted_connection=yes;')

connection = pyodbc.connect(connect_info)

cursor = connection.cursor()

#  using the manufacturer name and a list of upcs as parameters for cursor.execute()
manufacturer = 'Manufacturer_name_here'

upc_list = []

#  Creating modular placeholders for any arguments we need to input
args_1 = [manufacturer, *upc_list]
placeholders = ", ".join('?' * len(upc_list))

# Pulling option names from SQL server
opname_query =  """SELECT DISTINCT ps.optionname           
                   FROM table_view_1 AS t (NOLOCK)
                   LEFT JOIN table_view_1 AS tb (NOLOCK)
                   ON t.productid = tb.productid
                   WHERE t.manufacturer = ?
                   AND p.upc IN(%s);""" % placeholders

# Converting the pulled data to delimited values
cursor.execute(opname_query, args_1)
results_1 = cursor.fetchall()
results_1_df = pd.DataFrame(results_1, columns=[x[0] for x in cursor.description])

#  creating a dataframe for later use in a CTE from the opname_query
opname= []
for row, value in results_1_df.iteritems():
    for v in value:
        str = ''.join(v)
        new_values = '[%s],' % str
        opname.append(new_values)
        ser = pd.Series(opname)
        opname_df = pd.DataFrame()
        opname_df['optionname'] = ser.values
        value = opname_df.iloc[-1]['optionname'].replace(",","")
        opname_df.iloc[-1, opname_df.columns.get_loc('optionname')] = value

args_2 = [manufacturer, *upc_list]
placeholders_2 = ", ".join('?' * len(upc_list))

# Creating dataframe from the query results
specs_query =  """WITH SPECS AS(                                       
                    SELECT
                        p.upc,
                        p.manufacturer,
                        p.sku,
                        p.finish,
                        p.weight,
                        ps.optionname,
                        ps.optionvalue
                    FROM table_view_1 AS tb (NOLOCK)
                    INNER JOIN table_view_2 AS t (NOLOCK)
                    ON tb.productid = t.productid
                    WHERE p.manufacturer = ?
                    AND p.upc IN(%s))

                    SELECT *                                               
                    FROM SPECS""" % placeholders_2
                    
cursor.execute(specs_query, args_2)
results_2 = cursor.fetchall()
results_2_df = pd.DataFrame.from_records(results_2, columns=[x[0] for x in cursor.description])

#  Saving the DataFrame to a specified Excel workbook as a named worksheet
path = 'path\\to\\file\\Test_file.xlsx'
excel_wb = load_workbook(path)
with pd.ExcelWriter(path) as writer:
    writer.book = excel_wb
    results_2_df.to_excel(writer, sheet_name='DB_data')